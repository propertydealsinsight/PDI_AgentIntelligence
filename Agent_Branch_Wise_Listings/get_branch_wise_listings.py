#!/usr/bin/env python3
"""
Production utility: export raw branch listings + Birds Eye View for an agent brand.

1. Looks up distinct agent_master_id values from pdi_agent_performance by agent name.
2. Runs the per-branch listing query (unchanged) against pdi_agent_master for each id.
3. Writes incremental results to Excel, then adds a Birds Eye View summary tab.

Requirements (pip install -r requirements.txt):
  mysql-connector-python>=9.0.0
  openpyxl>=3.1.0
  pandas>=2.0.0

Setup:
  copy example.config.py to config.py and fill in credentials.

Usage:
  python get_branch_wise_listings.py --agent-name "Fine & Country"
  python get_branch_wise_listings.py --agent-name "Fine & Country%" --output output/my-run/Fine_and_Country.xlsx
  python get_branch_wise_listings.py --agent-name "Fine & Country" --select-branches
  python get_branch_wise_listings.py --agent-name "Belvoir" --resume
  nohup python get_branch_wise_listings.py --agent-name "Fine & Country" > export.log 2>&1 &
"""

from __future__ import annotations

import argparse
import csv
import logging
from logging.handlers import RotatingFileHandler
import os
import re
import signal
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import mysql.connector
import pandas as pd
from mysql.connector import Error as MySQLError
from mysql.connector.errors import InterfaceError, OperationalError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = SCRIPT_DIR / "output"
LOG_DIR = SCRIPT_DIR / "logs"
LOG_FILE = LOG_DIR / "get_branch_wise_listings.log"

LISTINGS_SHEET = "listings"
BIRDS_EYE_SHEET = "Birds Eye View"

SKIP_ON_RESUME = frozenset({"ok", "no_match"})
RETRY_ON_RESUME = frozenset({"error", "timeout", "connection_error", "excel_error", "unknown_error"})

AGENT_MASTER_IDS_SQL = """
SELECT
  p.agent_master_id,
  MAX(p.agent_name) AS agent_name,
  MAX(am.agent_name_rm) AS agent_name_rm,
  MAX(am.agent_name_zl) AS agent_name_zl,
  MAX(am.address_rm) AS address_rm,
  MAX(am.address_zl) AS address_zl
FROM PDI_PortalsData.pdi_agent_performance p
LEFT JOIN PDI_PortalsData.pdi_agent_master am ON am.id = p.agent_master_id
WHERE p.agent_name LIKE %s
  AND p.agent_master_id IS NOT NULL
GROUP BY p.agent_master_id
ORDER BY p.agent_master_id
"""

BRANCHES_CSV = "branches.csv"
BRANCHES_CSV_COLUMNS = [
    "agent_master_id",
    "agent_name",
    "agent_name_rm",
    "agent_name_zl",
    "address_rm",
    "address_zl",
]

# ---------------------------------------------------------------------------
# Listing query — unchanged from export_agent_raw_listings.py
# ---------------------------------------------------------------------------
LISTING_QUERY = """
WITH T_UnionTables AS (
    SELECT
        DISTINCT
        'Rightmove' AS portal,
        p.listing_id, p.property_identifier, agent_name, agent_address, p.agent_logo,
        first_published_date, listing_update_date, displayable_address, full_property_address,
        uprn, postcode, num_bedrooms,
        CASE
            WHEN LOWER(p.price_frequency) = 'weekly'    THEN ROUND(p.price * 52 / 12)
            WHEN LOWER(p.price_frequency) = 'monthly'   THEN p.price
            WHEN LOWER(p.price_frequency) = 'pppm'      THEN p.price
            WHEN LOWER(p.price_frequency) = 'quarterly' THEN ROUND(p.price / 3)
            WHEN LOWER(p.price_frequency) = 'yearly'    THEN ROUND(p.price / 12)
            ELSE p.price
        END AS price,
        CASE
            WHEN LOWER(p.price_frequency) = 'weekly'    THEN ROUND(p.listed_price * 52 / 12)
            WHEN LOWER(p.price_frequency) = 'monthly'   THEN p.listed_price
            WHEN LOWER(p.price_frequency) = 'pppm'      THEN p.listed_price
            WHEN LOWER(p.price_frequency) = 'quarterly' THEN ROUND(p.listed_price / 3)
            WHEN LOWER(p.price_frequency) = 'yearly'    THEN ROUND(p.listed_price / 12)
            ELSE p.listed_price
        END AS listed_price,
        COALESCE(p.pdi_ppd_last_transaction_soldPrice, p.last_transaction_soldPrice) AS sold_price,
        COALESCE(p.pdi_ppd_last_transaction_date, p.last_transaction_date)           AS txn_date,
        outcode, listing_status, display_status, status, listing_update_reason,
        CASE
            WHEN p.property_type IN ('Apartment') THEN 'Flat'
            WHEN p.property_type IN ('End of Terrace') THEN 'end_terrace'
            ELSE REPLACE(REPLACE(p.property_type, '-', '_'), ' ', '_')
        END AS property_type,
        details_url,
        (DATEDIFF(NOW(), first_published_date)) AS listed_days,
        IF(first_published_date >= DATE_SUB(CURDATE(), INTERVAL 365 DAY), 1, 0) AS is_recent,
        IF(p.display_status IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract'), 1, 0) AS marked_under_offer_sold,
        IF(IFNULL(p.display_status, '') NOT IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract') AND p.status IN ('for_sale', 'to_rent'), 1, 0) AS active_listing,
        IF(IFNULL(p.display_status, '') NOT IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract') AND p.status IN ('removed', 'archived', 'EXPIRED'), 1, 0) AS withdrawn_listing
    FROM PDI_PortalsData.property_details p
    JOIN PDI_PortalsData.pdi_agent_master pam ON (p.agent_name = pam.agent_name_rm AND p.agent_address = pam.address_rm)
    JOIN PDI_PortalsData.property_type_mapping p_mapping ON (p.property_type = p_mapping.property_type)
    WHERE pam.id = @agentId
      AND p_mapping.property_category IN ('flats', 'houses')
      AND p.listing_status IN ('sale', 'new-homes')
      AND p.first_published_date >= DATE_SUB(CURDATE(), INTERVAL 12 MONTH)
      AND p.residential = 'YES'
    UNION ALL
    SELECT
        DISTINCT
        'Zoopla' AS portal,
        p.listing_id, p.property_identifier, agent_name, agent_address, p.agent_logo,
        first_published_date, listing_update_date, displayable_address, full_property_address,
        uprn, postcode, num_bedrooms,
        CASE
            WHEN LOWER(p.price_frequency) = 'weekly'    THEN ROUND(p.price * 52 / 12)
            WHEN LOWER(p.price_frequency) = 'monthly'   THEN p.price
            WHEN LOWER(p.price_frequency) = 'pppm'      THEN p.price
            WHEN LOWER(p.price_frequency) = 'quarterly' THEN ROUND(p.price / 3)
            WHEN LOWER(p.price_frequency) = 'yearly'    THEN ROUND(p.price / 12)
            ELSE p.price
        END AS price,
        CASE
            WHEN LOWER(p.price_frequency) = 'weekly'    THEN ROUND(p.listed_price * 52 / 12)
            WHEN LOWER(p.price_frequency) = 'monthly'   THEN p.listed_price
            WHEN LOWER(p.price_frequency) = 'pppm'      THEN p.listed_price
            WHEN LOWER(p.price_frequency) = 'quarterly' THEN ROUND(p.listed_price / 3)
            WHEN LOWER(p.price_frequency) = 'yearly'    THEN ROUND(p.listed_price / 12)
            ELSE p.listed_price
        END AS listed_price,
        COALESCE(p.pdi_ppd_last_transaction_soldPrice, p.last_transaction_soldPrice) AS sold_price,
        COALESCE(p.pdi_ppd_last_transaction_date, p.last_transaction_date)           AS txn_date,
        outcode, listing_status, display_status, status, listing_update_reason,
        CASE
            WHEN p.property_type IN ('Apartment') THEN 'Flat'
            WHEN p.property_type IN ('End of Terrace') THEN 'end_terrace'
            ELSE REPLACE(REPLACE(p.property_type, '-', '_'), ' ', '_')
        END AS property_type,
        details_url,
        (DATEDIFF(NOW(), first_published_date)) AS listed_days,
        IF(first_published_date >= DATE_SUB(CURDATE(), INTERVAL 365 DAY), 1, 0) AS is_recent,
        IF(p.display_status IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract'), 1, 0) AS marked_under_offer_sold,
        IF(IFNULL(p.display_status, '') NOT IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract') AND p.status IN ('for_sale', 'to_rent'), 1, 0) AS active_listing,
        IF(IFNULL(p.display_status, '') NOT IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract') AND p.status IN ('removed', 'archived', 'EXPIRED'), 1, 0) AS withdrawn_listing
    FROM PDI_PortalsData.property_details_zoopla p
    JOIN PDI_PortalsData.pdi_agent_master pam ON (p.agent_name = pam.agent_name_zl AND p.agent_address = pam.address_zl)
    JOIN PDI_PortalsData.property_type_mapping p_mapping ON (p.property_type = p_mapping.property_type)
    WHERE pam.id = @agentId
      AND p_mapping.property_category IN ('flats', 'houses')
      AND p.listing_status IN ('sale', 'new-homes')
      AND p.first_published_date >= DATE_SUB(CURDATE(), INTERVAL 12 MONTH)
      AND p.category = 'residential'
),
T_rec_1 AS (
    SELECT t.*, ROW_NUMBER() OVER (PARTITION BY IF(t.listing_status='new-homes','sale',t.listing_status), COALESCE(t.uprn, LOWER(CONCAT(SUBSTRING_INDEX(REPLACE(REPLACE(REPLACE(UPPER(REPLACE(REPLACE(t.full_property_address, ', ', ' '), ' and parking space', '')), 'FLAT ', ''), 'APARTMENT ', ''), 'UNIT ', ''), ' ', 3), ' ', t.postcode)), LOWER(CONCAT_WS('|', t.postcode, t.property_type, t.num_bedrooms, t.price))), IF(t.listed_days < 4, 0, t.listed_days) ORDER BY t.marked_under_offer_sold DESC, t.first_published_date DESC) AS rn1
    FROM (SELECT * FROM T_UnionTables WHERE is_recent = 1) t
),
T_rec_2 AS (
    SELECT t.*, ROW_NUMBER() OVER (PARTITION BY IF(t.listing_status='new-homes','sale',t.listing_status), LOWER(CONCAT_WS('|', t.postcode, t.property_type, t.num_bedrooms, t.price)), IF(t.listed_days < 4, 0, t.listed_days) ORDER BY t.marked_under_offer_sold DESC, t.first_published_date DESC) AS rn2
    FROM T_rec_1 t WHERE t.rn1 = 1
),
T_rec_final AS (SELECT * FROM T_rec_2 WHERE rn2 = 1),
T_all_1 AS (
    SELECT t.*, ROW_NUMBER() OVER (PARTITION BY IF(t.listing_status='new-homes','sale',t.listing_status), COALESCE(t.uprn, LOWER(CONCAT(SUBSTRING_INDEX(REPLACE(REPLACE(REPLACE(UPPER(REPLACE(REPLACE(t.full_property_address, ', ', ' '), ' and parking space', '')), 'FLAT ', ''), 'APARTMENT ', ''), 'UNIT ', ''), ' ', 3), ' ', t.postcode)), LOWER(CONCAT_WS('|', t.postcode, t.property_type, t.num_bedrooms, t.price))), IF(t.listed_days < 4, 0, t.listed_days) ORDER BY t.marked_under_offer_sold DESC, t.first_published_date DESC) AS rn1
    FROM T_UnionTables t
),
T_all_2 AS (
    SELECT t.*, ROW_NUMBER() OVER (PARTITION BY IF(t.listing_status='new-homes','sale',t.listing_status), LOWER(CONCAT_WS('|', t.postcode, t.property_type, t.num_bedrooms, t.price)), IF(t.listed_days < 4, 0, t.listed_days) ORDER BY t.marked_under_offer_sold DESC, t.first_published_date DESC) AS rn2
    FROM T_all_1 t WHERE t.rn1 = 1
),
T_all_final AS (SELECT * FROM T_all_2 WHERE rn2 = 1),
T_sold AS (
    SELECT
        t.listing_id,
        CASE
            WHEN t.listed_price IS NOT NULL AND t.listed_price > 1000 THEN t.listed_price
            WHEN t.price        IS NOT NULL AND t.price        > 1000 THEN t.price
            ELSE NULL
        END AS asking,
        IF((t.listed_price IS NULL OR t.listed_price <= 1000)
           AND t.price > 1000, 1, 0) AS asking_recovered,
        t.sold_price, t.txn_date, t.first_published_date
    FROM T_all_final t
    WHERE t.marked_under_offer_sold = 1
),
T_valid AS (
    SELECT listing_id, asking_recovered,
           ROUND(((sold_price - asking) / asking) * 100, 2) AS variation_pct
    FROM T_sold
    WHERE asking IS NOT NULL AND asking > 0
      AND sold_price IS NOT NULL AND sold_price > 0
      AND txn_date IS NOT NULL
      AND txn_date >= first_published_date
),
T_med_raw AS (
    SELECT variation_pct,
           ROW_NUMBER() OVER (ORDER BY variation_pct) AS rn,
           COUNT(*)     OVER ()                       AS cnt
    FROM T_valid
),
T_med_raw_val AS (
    SELECT AVG(variation_pct) AS med, MAX(cnt) AS n_valid
    FROM T_med_raw
    WHERE rn IN (FLOOR((cnt + 1) / 2), CEIL((cnt + 1) / 2))
),
T_mad AS (
    SELECT ABS(v.variation_pct - r.med)                                   AS abs_dev,
           ROW_NUMBER() OVER (ORDER BY ABS(v.variation_pct - r.med))      AS rn,
           COUNT(*)     OVER ()                                           AS cnt
    FROM T_valid v CROSS JOIN T_med_raw_val r
),
T_mad_val AS (
    SELECT AVG(abs_dev) AS mad
    FROM T_mad
    WHERE rn IN (FLOOR((cnt + 1) / 2), CEIL((cnt + 1) / 2))
),
T_fence AS (
    SELECT
        r.n_valid,
        CASE WHEN r.n_valid >= 5 AND m.mad > 0
             THEN r.med - 3.5 * 1.4826 * m.mad END AS lo,
        CASE WHEN r.n_valid >= 5 AND m.mad > 0
             THEN r.med + 3.5 * 1.4826 * m.mad END AS hi
    FROM T_med_raw_val r CROSS JOIN T_mad_val m
),
T_kept AS (
    SELECT v.listing_id, v.variation_pct
    FROM T_valid v CROSS JOIN T_fence f
    WHERE f.lo IS NULL OR v.variation_pct BETWEEN f.lo AND f.hi
),
T_ranked AS (
    SELECT variation_pct,
           ROW_NUMBER() OVER (ORDER BY variation_pct) AS rn,
           COUNT(*)     OVER ()                       AS cnt
    FROM T_kept
),
T_turnaround AS (
    SELECT days_diff FROM (
        SELECT
            DATEDIFF(ph.changed_at,
                IF(t.first_published_date > t.listing_update_date, t.listing_update_date, t.first_published_date)) AS days_diff,
            ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS rns
        FROM T_rec_final t
        JOIN PDI_PortalsData.property_details_history ph
            ON ph.column_name = 'display_status' AND ph.listing_id = t.listing_id
        WHERE t.portal = 'Rightmove' AND t.marked_under_offer_sold = 1
        UNION ALL
        SELECT
            DATEDIFF(ph.changed_at,
                IF(t.first_published_date > t.listing_update_date, t.listing_update_date, t.first_published_date)) AS days_diff,
            ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS rns
        FROM T_rec_final t
        JOIN PDI_PortalsData.property_details_zoopla_history ph
            ON ph.column_name = 'display_status' AND ph.listing_id = t.listing_id
        WHERE t.portal = 'Zoopla' AND t.marked_under_offer_sold = 1
    ) x
    WHERE x.rns = 1 AND x.days_diff > 1
),
T_rm_mark_sold_stc_history_date AS (
    Select t.*, ph.changed_at as marked_sold_stc_date, ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS row_num_marked_sold_stc
    from T_rec_final t JOIN
    PDI_PortalsData.property_details_history ph ON
    (t.listing_id = ph.listing_id and ph.column_name = 'display_status' and ph.new_value IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract'))
    where t.portal = 'Rightmove'
    and t.marked_under_offer_sold = 1
),
T_zl_mark_sold_stc_history_date AS (
    Select t.*, ph.changed_at as marked_sold_stc_date, ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS row_num_marked_sold_stc
    from T_rec_final t JOIN
    PDI_PortalsData.property_details_zoopla_history ph ON
    (t.listing_id = ph.listing_id and ph.column_name = 'display_status' and ph.new_value IN ('Sold STC', 'Under offer', 'Reserved', 'Sold STCM', 'Sold subject to', 'Under offer', 'Sold subject to contract'))
    where t.portal = 'Zoopla'
    and t.marked_under_offer_sold = 1
),
T_rm_mark_withdrawn_date AS (
    Select t.*, ph.changed_at as withdrawn_date, ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS row_num_withdrawn_dt
    from T_rec_final t JOIN
    PDI_PortalsData.property_details_history ph ON
    (t.listing_id = ph.listing_id and ph.column_name = 'status' and ph.new_value IN ('EXPIRED', 'removed', 'archived'))
    where t.portal = 'Rightmove'
    and t.withdrawn_listing = 1
),
T_zl_mark_withdrawn_history_date AS (
    Select t.*, ph.changed_at as withdrawn_date, ROW_NUMBER() OVER (PARTITION BY t.listing_id ORDER BY ph.changed_at DESC) AS row_num_withdrawn_dt
    from T_rec_final t JOIN
    PDI_PortalsData.property_details_zoopla_history ph ON
    (t.listing_id = ph.listing_id and ph.column_name = 'status' and ph.new_value IN ('EXPIRED', 'removed', 'archived'))
    where t.portal = 'Zoopla'
    and t.withdrawn_listing = 1
),
T_price_paid_evidence As (
    Select t.*, pam.lastSoldDate, pam.lastSoldPrice from
    T_rec_final t JOIN PDI_PortalsData.pdi_address_master pam ON
    (t.uprn is not null and t.uprn = pam.uprn and pam.lastSoldDate is not null and pam.lastSoldDate > t.first_published_date)
    WHERE t.marked_under_offer_sold = 1
)
SELECT
    t.portal,
    t.listing_id,
    t.property_identifier,
    t.agent_name,
    t.agent_address,
    t.agent_logo,
    t.first_published_date,
    t.listing_update_date,
    t.displayable_address,
    t.full_property_address,
    t.uprn,
    t.postcode,
    t.num_bedrooms,
    t.price,
    t.listed_price,
    t.sold_price,
    t.txn_date,
    t.outcode,
    t.listing_status,
    t.display_status,
    t.status,
    t.listing_update_reason,
    t.property_type,
    t.details_url,
    t.listed_days,
    t.is_recent,
    t.marked_under_offer_sold,
    t.active_listing,
    t.withdrawn_listing,
    COALESCE(t_rm_sold_dt.marked_sold_stc_date, t_zl_sold_dt.marked_sold_stc_date) as marked_sold_stc_under_offer_date,
    COALESCE(t_rm_withdrawn_dt.withdrawn_date, t_zl_withdrawn_dt.withdrawn_date) as withdrawn_date,
    t_price_paid.lastSoldDate as lastSoldDateEvidence,
    t_price_paid.lastSoldPrice as lastSoldPriceEvidence
FROM T_rec_final t
LEFT JOIN T_rm_mark_sold_stc_history_date t_rm_sold_dt ON
    (t.listing_id = t_rm_sold_dt.listing_id and t_rm_sold_dt.row_num_marked_sold_stc = 1)
LEFT JOIN T_zl_mark_sold_stc_history_date t_zl_sold_dt ON
    (t.listing_id = t_zl_sold_dt.listing_id and t_zl_sold_dt.row_num_marked_sold_stc = 1)
LEFT JOIN T_rm_mark_withdrawn_date t_rm_withdrawn_dt ON
    (t.listing_id = t_rm_withdrawn_dt.listing_id and t_rm_withdrawn_dt.row_num_withdrawn_dt = 1)
LEFT JOIN T_zl_mark_withdrawn_history_date t_zl_withdrawn_dt ON
    (t.listing_id = t_zl_withdrawn_dt.listing_id and t_zl_withdrawn_dt.row_num_withdrawn_dt = 1)
LEFT JOIN T_price_paid_evidence t_price_paid ON
    (t.listing_id = t_price_paid.listing_id)
"""

RESULT_COLUMNS = [
    "agent_master_id",
    "portal",
    "listing_id",
    "property_identifier",
    "agent_name",
    "agent_address",
    "agent_logo",
    "first_published_date",
    "listing_update_date",
    "displayable_address",
    "full_property_address",
    "uprn",
    "postcode",
    "num_bedrooms",
    "price",
    "listed_price",
    "sold_price",
    "txn_date",
    "outcode",
    "listing_status",
    "display_status",
    "status",
    "listing_update_reason",
    "property_type",
    "details_url",
    "listed_days",
    "is_recent",
    "marked_under_offer_sold",
    "active_listing",
    "withdrawn_listing",
    "marked_sold_stc_under_offer_date",
    "withdrawn_date",
    "lastSoldDateEvidence",
    "lastSoldPriceEvidence",
]

def setup_logging() -> logging.Logger:
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    log_format = "%(asctime)s %(levelname)s %(message)s"
    date_format = "%Y-%m-%d %H:%M:%S"
    formatter = logging.Formatter(log_format, datefmt=date_format)

    root = logging.getLogger()
    if root.handlers:
        return logging.getLogger(__name__)

    root.setLevel(logging.INFO)

    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(formatter)
    root.addHandler(console)

    file_handler = RotatingFileHandler(
        LOG_FILE,
        maxBytes=5 * 1024 * 1024,
        backupCount=3,
        encoding="utf-8",
    )
    file_handler.setFormatter(formatter)
    root.addHandler(file_handler)

    return logging.getLogger(__name__)


logger = setup_logging()

_shutdown_requested = False


def register_signal_handlers() -> None:
    def _handler(signum: int, _frame: Any) -> None:
        global _shutdown_requested
        _shutdown_requested = True
        logger.warning("Shutdown requested — will stop after current branch id.")

    signal.signal(signal.SIGINT, _handler)
    signal.signal(signal.SIGTERM, _handler)
    if hasattr(signal, "SIGBREAK"):
        signal.signal(signal.SIGBREAK, _handler)


def load_db_config() -> dict[str, Any]:
    cfg_host = ""
    cfg_user = ""
    cfg_pass = ""
    cfg_db = "PDI_PortalsData"
    cfg_port = 3306

    sys.path.insert(0, str(SCRIPT_DIR))
    try:
        import config as cfg  # type: ignore

        cfg_host = cfg.HOST
        cfg_user = cfg.USER
        cfg_pass = cfg.PASSWORD
        cfg_db = getattr(cfg, "DATABASE", cfg_db)
        cfg_port = getattr(cfg, "PORT", cfg_port)
    except ModuleNotFoundError as exc:
        raise SystemExit("config.py not found. Copy example.config.py to config.py first.") from exc

    return {
        "host": os.environ.get("PDI_RO_HOST", cfg_host),
        "port": int(os.environ.get("PDI_RO_PORT", str(cfg_port))),
        "user": os.environ.get("PDI_RO_USER", cfg_user),
        "password": os.environ.get("PDI_RO_PASSWORD", cfg_pass),
        "database": os.environ.get("PDI_RO_DB", cfg_db),
    }


def connect_db(cfg: dict[str, Any], *, read_timeout: int) -> mysql.connector.MySQLConnection:
    if not cfg["password"]:
        raise SystemExit("Database password not set in config.py")
    return mysql.connector.connect(
        host=cfg["host"],
        port=cfg["port"],
        user=cfg["user"],
        password=cfg["password"],
        database=cfg["database"],
        connection_timeout=60,
        read_timeout=read_timeout,
        write_timeout=30,
        autocommit=True,
    )


def reconnect_db(
    cfg: dict[str, Any],
    conn: mysql.connector.MySQLConnection | None,
    *,
    read_timeout: int,
) -> mysql.connector.MySQLConnection:
    if conn is not None:
        try:
            conn.close()
        except Exception:
            pass
    logger.info("Reconnecting to database...")
    return connect_db(cfg, read_timeout=read_timeout)


def like_pattern(agent_name: str) -> str:
    if "%" in agent_name or "_" in agent_name:
        return agent_name
    return f"{agent_name}%"


def safe_filename(agent_name: str) -> str:
    cleaned = agent_name.replace("%", "").strip()
    cleaned = re.sub(r"[^\w\s-]", "", cleaned)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("_")
    return cleaned or "agent_export"


def listings_filename(agent_name: str) -> str:
    return f"{safe_filename(agent_name)}_raw_listings.xlsx"


def run_dir_name(agent_name: str) -> str:
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    return f"{stamp}_{safe_filename(agent_name)}"


def create_run_output_dir(agent_name: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    while True:
        run_dir = OUTPUT_DIR / run_dir_name(agent_name)
        if not run_dir.exists():
            run_dir.mkdir(parents=True)
            return run_dir
        time.sleep(1)


def find_latest_run_output(agent_name: str) -> Path | None:
    agent_slug = safe_filename(agent_name)
    suffix = f"_{agent_slug}"
    run_dirs = sorted(
        (
            path
            for path in OUTPUT_DIR.iterdir()
            if path.is_dir() and path.name.endswith(suffix)
        ),
        key=lambda path: path.name,
        reverse=True,
    )
    for run_dir in run_dirs:
        output = run_dir / listings_filename(agent_name)
        if output.is_file():
            return output
    return None


def resolve_output_path(
    agent_name: str,
    explicit: Path | None,
    *,
    resume: bool,
) -> Path:
    if explicit is not None:
        explicit = explicit.resolve()
        if explicit.suffix.lower() == ".xlsx":
            explicit.parent.mkdir(parents=True, exist_ok=True)
            return explicit
        explicit.mkdir(parents=True, exist_ok=True)
        return explicit / listings_filename(agent_name)

    if resume:
        latest = find_latest_run_output(agent_name)
        if latest is not None:
            return latest

    run_dir = create_run_output_dir(agent_name)
    return run_dir / listings_filename(agent_name)


def progress_path(output: Path) -> Path:
    return output.with_suffix(output.suffix + ".progress")


def errors_csv_path(output: Path) -> Path:
    return output.with_suffix(output.suffix + ".errors.csv")


def failed_ids_path(output: Path) -> Path:
    return output.with_suffix(output.suffix + ".failed_ids.txt")


def branches_csv_path(output: Path) -> Path:
    return output.parent / BRANCHES_CSV


def write_branches_csv(csv_path: Path, branches: list[dict[str, Any]]) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=BRANCHES_CSV_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for branch in branches:
            writer.writerow({col: branch.get(col, "") for col in BRANCHES_CSV_COLUMNS})


def read_branch_ids_from_csv(csv_path: Path) -> list[int]:
    if not csv_path.is_file():
        raise FileNotFoundError(f"Branch list not found: {csv_path}")

    branch_ids: list[int] = []
    seen: set[int] = set()
    with csv_path.open(encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        if not reader.fieldnames or "agent_master_id" not in reader.fieldnames:
            raise ValueError(f"{csv_path.name} must include an agent_master_id column.")
        for row_num, row in enumerate(reader, start=2):
            raw = (row.get("agent_master_id") or "").strip()
            if not raw:
                continue
            try:
                branch_id = int(raw)
            except ValueError as exc:
                raise ValueError(
                    f"Invalid agent_master_id on line {row_num} in {csv_path.name}: {raw!r}"
                ) from exc
            if branch_id not in seen:
                seen.add(branch_id)
                branch_ids.append(branch_id)
    return branch_ids


def confirm_branch_selection(csv_path: Path) -> list[int]:
    print()
    print("Review the branch list before export:")
    print(f"  {csv_path}")
    print()
    print("Remove rows for branches you do NOT want to export, save the file,")
    print("then press Y to continue or N to cancel.")
    print()

    while True:
        try:
            answer = input("Continue? [Y/N]: ").strip().upper()
        except EOFError as exc:
            raise SystemExit(
                "Branch selection requires interactive input. Re-run from a terminal."
            ) from exc

        if answer == "N":
            raise SystemExit("Cancelled by user.")
        if answer == "Y":
            branch_ids = read_branch_ids_from_csv(csv_path)
            if not branch_ids:
                print("No branches found in the file. Keep at least one row or press N to cancel.")
                continue
            print(f"Processing {len(branch_ids)} branch(es).")
            return branch_ids
        print("Please enter Y or N.")


def resolve_agent_ids(
    branches: list[dict[str, Any]],
    output: Path,
    *,
    select_branches: bool,
    resume: bool,
) -> list[int]:
    if not select_branches:
        return [int(branch["agent_master_id"]) for branch in branches]

    csv_path = branches_csv_path(output)
    if resume and csv_path.is_file():
        branch_ids = read_branch_ids_from_csv(csv_path)
        logger.info("Loaded %s branch(es) from %s", len(branch_ids), csv_path)
        return branch_ids

    write_branches_csv(csv_path, branches)
    logger.info("Wrote %s branch(es) to %s for review", len(branches), csv_path)
    return confirm_branch_selection(csv_path)


def parse_progress_status(raw_status: str) -> str:
    status = raw_status.strip().lower()
    if status in SKIP_ON_RESUME or status in RETRY_ON_RESUME:
        return status
    if status.startswith("error"):
        return "error"
    return status


def load_processed_ids(
    output: Path,
    *,
    resume: bool,
    skip_failed: bool,
) -> set[int]:
    processed: set[int] = set()
    latest_status: dict[int, str] = {}

    progress_file = progress_path(output)
    if resume and progress_file.exists():
        for line in progress_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split(",", 2)
            if len(parts) < 2:
                continue
            agent_id = int(parts[0])
            status = parse_progress_status(parts[1])
            latest_status[agent_id] = status
            if status in SKIP_ON_RESUME:
                processed.add(agent_id)
            elif status in RETRY_ON_RESUME and skip_failed:
                processed.add(agent_id)

    if resume and output.exists():
        try:
            wb = load_workbook(output, read_only=True, data_only=True)
            ws = wb[LISTINGS_SHEET] if LISTINGS_SHEET in wb.sheetnames else wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    processed.add(int(row[0]))
            wb.close()
        except Exception as exc:
            logger.warning("Could not read existing Excel for resume: %s", exc)

    return processed


def append_progress(output: Path, agent_id: int, status: str, row_count: int = 0) -> None:
    with progress_path(output).open("a", encoding="utf-8") as fh:
        fh.write(
            f"{agent_id},{status},{row_count},"
            f"{datetime.now().isoformat(timespec='seconds')}\n"
        )


class ErrorLogger:
    CSV_HEADER = [
        "agent_master_id",
        "error_type",
        "mysql_errno",
        "message",
        "duration_sec",
        "attempts",
        "timestamp",
    ]

    def __init__(self, output: Path) -> None:
        self.csv_path = errors_csv_path(output)
        self.ids_path = failed_ids_path(output)
        if not self.csv_path.exists():
            with self.csv_path.open("w", encoding="utf-8", newline="") as fh:
                csv.writer(fh).writerow(self.CSV_HEADER)

    def log_failure(
        self,
        agent_id: int,
        error_type: str,
        message: str,
        *,
        mysql_errno: int | None = None,
        duration_sec: float = 0.0,
        attempts: int = 1,
    ) -> None:
        with self.csv_path.open("a", encoding="utf-8", newline="") as fh:
            csv.writer(fh).writerow(
                [
                    agent_id,
                    error_type,
                    mysql_errno if mysql_errno is not None else "",
                    message.replace("\n", " ")[:2000],
                    round(duration_sec, 2),
                    attempts,
                    datetime.now().isoformat(timespec="seconds"),
                ]
            )
        with self.ids_path.open("a", encoding="utf-8") as fh:
            fh.write(f"{agent_id}\n")


def classify_mysql_error(exc: MySQLError) -> str:
    errno = getattr(exc, "errno", None)
    if errno in (3024, 1317, 1205):
        return "timeout"
    if errno in (2006, 2013):
        return "connection_error"
    msg = str(exc).lower()
    if "timeout" in msg or "timed out" in msg:
        return "timeout"
    if "lost connection" in msg or "gone away" in msg:
        return "connection_error"
    return "error"


def cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def append_excel_rows(output: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    if output.exists():
        wb = load_workbook(output)
        ws = wb[LISTINGS_SHEET] if LISTINGS_SHEET in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = LISTINGS_SHEET
        ws.append(columns)

    for row in rows:
        ws.append([cell_value(row.get(col)) for col in columns])

    wb.save(output)
    wb.close()


def set_session_query_timeout(cursor: Any, timeout_sec: int) -> None:
    cursor.execute(f"SET SESSION MAX_EXECUTION_TIME = {max(timeout_sec, 1) * 1000}")


def fetch_agent_branches(
    conn: mysql.connector.MySQLConnection,
    like_value: str,
) -> list[dict[str, Any]]:
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(AGENT_MASTER_IDS_SQL, (like_value,))
        return cursor.fetchall()
    finally:
        cursor.close()


def fetch_listings(
    conn: mysql.connector.MySQLConnection,
    agent_id: int,
    *,
    query_timeout_sec: int,
    max_retries: int,
) -> tuple[list[dict[str, Any]], int]:
    last_error: Exception | None = None
    last_error_type = "error"
    attempts_used = 0

    for attempt in range(1, max_retries + 1):
        attempts_used = attempt
        cursor = conn.cursor(dictionary=True)
        try:
            set_session_query_timeout(cursor, query_timeout_sec)
            cursor.execute("SET @agentId := %s", (agent_id,))
            cursor.execute(LISTING_QUERY)
            rows = cursor.fetchall()
            for row in rows:
                row["agent_master_id"] = agent_id
            return rows, attempts_used
        except (MySQLError, InterfaceError, OperationalError) as exc:
            last_error = exc
            last_error_type = classify_mysql_error(exc) if isinstance(exc, MySQLError) else "connection_error"
            logger.warning(
                "Query failed for id=%s (attempt %s/%s, type=%s): %s",
                agent_id,
                attempt,
                max_retries,
                last_error_type,
                exc,
            )
            if last_error_type == "timeout":
                break
            if attempt < max_retries and last_error_type in {"connection_error", "error"}:
                time.sleep(min(2 ** attempt, 15))
        finally:
            cursor.close()

    assert last_error is not None
    err = RuntimeError(str(last_error))
    err.error_type = last_error_type  # type: ignore[attr-defined]
    err.mysql_errno = getattr(last_error, "errno", None) if isinstance(last_error, MySQLError) else None  # type: ignore[attr-defined]
    err.attempts = attempts_used  # type: ignore[attr-defined]
    raise err from last_error


# ---------------------------------------------------------------------------
# Birds Eye View — unchanged from add_birds_eye_summary.py
# ---------------------------------------------------------------------------
def effective_asking(row: pd.Series) -> float | None:
    listed = row.get("listed_price")
    price = row.get("price")
    if pd.notna(listed) and listed > 1000:
        return float(listed)
    if pd.notna(price) and price > 1000:
        return float(price)
    return None


def portal_first(group: pd.DataFrame, portal: str, column: str) -> str | None:
    subset = group[group["portal"] == portal][column].dropna()
    if subset.empty:
        return None
    return str(subset.iloc[0])


def portal_label(group: pd.DataFrame) -> str:
    portals = set(group["portal"].dropna().unique())
    if portals == {"Rightmove"}:
        return "Rightmove"
    if portals == {"Zoopla"}:
        return "Zoopla"
    if portals:
        return "Both"
    return ""


def branch_display_name(group: pd.DataFrame) -> str:
    return (
        portal_first(group, "Rightmove", "agent_name")
        or portal_first(group, "Zoopla", "agent_name")
        or ""
    )


def branch_display_address(group: pd.DataFrame) -> str:
    return (
        portal_first(group, "Rightmove", "agent_address")
        or portal_first(group, "Zoopla", "agent_address")
        or ""
    )


def build_branch_summary(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    work["effective_asking"] = work.apply(effective_asking, axis=1)
    work["first_published_date"] = pd.to_datetime(work["first_published_date"], errors="coerce")
    work["marked_sold_stc_under_offer_date"] = pd.to_datetime(
        work["marked_sold_stc_under_offer_date"], errors="coerce"
    )
    work["withdrawn_date"] = pd.to_datetime(work["withdrawn_date"], errors="coerce")

    rows: list[dict[str, Any]] = []

    for branch_id, g in work.groupby("agent_master_id", sort=True):
        total = len(g)
        active = int(g["active_listing"].sum())
        sold_uo = int(g["marked_under_offer_sold"].sum())
        withdrawn = int(g["withdrawn_listing"].sum())

        if active + sold_uo + withdrawn != total:
            raise ValueError(
                f"Branch {branch_id}: active+sold+withdrawn={active+sold_uo+withdrawn} "
                f"!= total={total}"
            )

        active_rows = g[g["active_listing"] == 1]
        sold_rows = g[g["marked_under_offer_sold"] == 1]

        rows.append(
            {
                "Branch ID": int(branch_id),
                "Agent Name": branch_display_name(g),
                "Branch Address": branch_display_address(g),
                "Portals": portal_label(g),
                "Total Listings": total,
                "Active": active,
                "Sold / Under Offer": sold_uo,
                "Withdrawn": withdrawn,
                "Active %": round(active / total * 100, 1) if total else 0,
                "Sold / UO %": round(sold_uo / total * 100, 1) if total else 0,
                "Withdrawn %": round(withdrawn / total * 100, 1) if total else 0,
                "Recent (365d window)": int(g["is_recent"].sum()),
                "Rightmove Listings": int((g["portal"] == "Rightmove").sum()),
                "Zoopla Listings": int((g["portal"] == "Zoopla").sum()),
                "Median Asking Price": round(g["effective_asking"].median(), 0)
                if g["effective_asking"].notna().any()
                else None,
                "Median Listed Days": round(g["listed_days"].median(), 0)
                if g["listed_days"].notna().any()
                else None,
                "Avg Listed Days (Active)": round(active_rows["listed_days"].mean(), 1)
                if len(active_rows)
                else None,
                "Avg Listed Days (Sold/UO)": round(sold_rows["listed_days"].mean(), 1)
                if len(sold_rows)
                else None,
                "Sold w/ Land Registry Price": int(sold_rows["sold_price"].notna().sum()),
                "Sold w/ Price-Paid Evidence": int(g["lastSoldDateEvidence"].notna().sum()),
                "Distinct Postcodes": int(g["postcode"].nunique(dropna=True)),
                "Distinct Outcodes": int(g["outcode"].nunique(dropna=True)),
                "Sale Listings": int((g["listing_status"] == "sale").sum()),
                "New Homes Listings": int((g["listing_status"] == "new-homes").sum()),
                "Earliest Published": g["first_published_date"].min(),
                "Latest Published": g["first_published_date"].max(),
            }
        )

    summary = pd.DataFrame(rows)
    branch_count = len(summary)
    total_row = {
        "Branch ID": "TOTAL",
        "Agent Name": f"{branch_count} branches",
        "Branch Address": "",
        "Portals": "",
        "Total Listings": int(summary["Total Listings"].sum()),
        "Active": int(summary["Active"].sum()),
        "Sold / Under Offer": int(summary["Sold / Under Offer"].sum()),
        "Withdrawn": int(summary["Withdrawn"].sum()),
        "Active %": round(summary["Active"].sum() / summary["Total Listings"].sum() * 100, 1),
        "Sold / UO %": round(
            summary["Sold / Under Offer"].sum() / summary["Total Listings"].sum() * 100, 1
        ),
        "Withdrawn %": round(
            summary["Withdrawn"].sum() / summary["Total Listings"].sum() * 100, 1
        ),
        "Recent (365d window)": int(summary["Recent (365d window)"].sum()),
        "Rightmove Listings": int(summary["Rightmove Listings"].sum()),
        "Zoopla Listings": int(summary["Zoopla Listings"].sum()),
        "Median Asking Price": round(work["effective_asking"].median(), 0)
        if work["effective_asking"].notna().any()
        else None,
        "Median Listed Days": round(work["listed_days"].median(), 0)
        if work["listed_days"].notna().any()
        else None,
        "Avg Listed Days (Active)": round(
            work.loc[work["active_listing"] == 1, "listed_days"].mean(), 1
        )
        if (work["active_listing"] == 1).any()
        else None,
        "Avg Listed Days (Sold/UO)": round(
            work.loc[work["marked_under_offer_sold"] == 1, "listed_days"].mean(), 1
        )
        if (work["marked_under_offer_sold"] == 1).any()
        else None,
        "Sold w/ Land Registry Price": int(
            work.loc[work["marked_under_offer_sold"] == 1, "sold_price"].notna().sum()
        ),
        "Sold w/ Price-Paid Evidence": int(work["lastSoldDateEvidence"].notna().sum()),
        "Distinct Postcodes": int(work["postcode"].nunique(dropna=True)),
        "Distinct Outcodes": int(work["outcode"].nunique(dropna=True)),
        "Sale Listings": int((work["listing_status"] == "sale").sum()),
        "New Homes Listings": int((work["listing_status"] == "new-homes").sum()),
        "Earliest Published": work["first_published_date"].min(),
        "Latest Published": work["first_published_date"].max(),
    }
    return pd.concat([summary, pd.DataFrame([total_row])], ignore_index=True)


def style_birds_eye_sheet(path: Path, n_rows: int, n_cols: int) -> None:
    wb = load_workbook(path)
    ws = wb[BIRDS_EYE_SHEET]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    total_fill = PatternFill("solid", fgColor="D9E1F2")
    total_font = Font(bold=True)

    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell = ws.cell(row=n_rows, column=col)
        cell.fill = total_fill
        cell.font = total_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows - 1}"

    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        max_len = max(
            (len(str(ws[f"{letter}{row}"].value)) for row in range(1, n_rows + 1) if ws[f"{letter}{row}"].value is not None),
            default=0,
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)

    wb.save(path)
    wb.close()


def add_birds_eye_sheet(output: Path) -> pd.DataFrame:
    df = pd.read_excel(output, sheet_name=LISTINGS_SHEET)
    if df.empty:
        raise ValueError("No listing rows found — Birds Eye View not created.")

    summary = build_branch_summary(df)
    with pd.ExcelWriter(
        output,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        summary.to_excel(writer, sheet_name=BIRDS_EYE_SHEET, index=False)

    style_birds_eye_sheet(output, n_rows=len(summary) + 1, n_cols=len(summary.columns))
    return summary


def log_failure(
    output: Path,
    error_logger: ErrorLogger,
    agent_id: int,
    *,
    error_type: str,
    message: str,
    duration_sec: float,
    mysql_errno: int | None = None,
    attempts: int = 1,
) -> None:
    append_progress(output, agent_id, error_type, 0)
    error_logger.log_failure(
        agent_id,
        error_type,
        message,
        mysql_errno=mysql_errno,
        duration_sec=duration_sec,
        attempts=attempts,
    )
    logger.error(
        "id=%s FAILED (%s, %.1fs): %s — logged to %s",
        agent_id,
        error_type,
        duration_sec,
        message,
        error_logger.csv_path.name,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export raw listings + Birds Eye View for an agent brand by name."
    )
    parser.add_argument(
        "--agent-name",
        required=True,
        help='Agent name search pattern, e.g. "Fine & Country" or "Fine & Country%%"',
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help=(
            "Excel file or output directory (default: "
            "output/<timestamp>_<AgentName>/<AgentName>_raw_listings.xlsx)"
        ),
    )
    parser.add_argument(
        "--select-branches",
        action="store_true",
        help=(
            "Write branches.csv for manual review, wait for Y, then export only "
            "the branches remaining in that file"
        ),
    )
    parser.add_argument("--resume", action="store_true", help="Skip completed branch ids")
    parser.add_argument(
        "--skip-failed",
        action="store_true",
        help="With --resume, do not retry ids that previously failed",
    )
    parser.add_argument("--limit", type=int, default=0, help="Process only first N pending ids")
    parser.add_argument("--max-retries", type=int, default=3)
    parser.add_argument("--query-timeout", type=int, default=180, help="Per-branch query timeout (seconds)")
    parser.add_argument("--sleep", type=float, default=0.0, help="Pause between branch ids")
    parser.add_argument(
        "--skip-birds-eye",
        action="store_true",
        help="Skip Birds Eye View tab (listings export only)",
    )
    return parser.parse_args()


def main() -> int:
    global _shutdown_requested

    args = parse_args()
    register_signal_handlers()

    cfg = load_db_config()

    pattern = like_pattern(args.agent_name)
    output = resolve_output_path(args.agent_name, args.output, resume=args.resume)

    read_timeout = args.query_timeout + 60
    error_logger = ErrorLogger(output)

    logger.info("Log file: %s", LOG_FILE)
    logger.info("Output folder: %s", output.parent)
    logger.info("Output file: %s", output)
    logger.info("Connecting to %s as %s", cfg["host"], cfg["user"])
    conn = connect_db(cfg, read_timeout=read_timeout)

    exit_code = 0
    total_rows = 0
    errors = 0
    interrupted = False

    try:
        branches = fetch_agent_branches(conn, pattern)
        if not branches:
            logger.error('No agent_master_id found for agent_name LIKE "%s"', pattern)
            return 1

        agent_ids = resolve_agent_ids(
            branches,
            output,
            select_branches=args.select_branches,
            resume=args.resume,
        )
        if not agent_ids:
            logger.error("No branches selected for export.")
            return 1

        processed = load_processed_ids(output, resume=args.resume, skip_failed=args.skip_failed)
        pending = [aid for aid in agent_ids if aid not in processed]
        if args.limit > 0:
            pending = pending[: args.limit]

        logger.info(
            'Agent name LIKE "%s" | branches found=%s | pending=%s | output=%s',
            pattern,
            len(agent_ids),
            len(pending),
            output,
        )
        logger.info("Error log: %s", error_logger.csv_path)

        if not pending and not output.exists():
            logger.info("Nothing to do.")
            return 0

        for idx, agent_id in enumerate(pending, start=1):
            if _shutdown_requested:
                interrupted = True
                logger.warning("Stopping before id=%s. Re-run with --resume to continue.", agent_id)
                break

            started = time.perf_counter()
            logger.info("[%s/%s] Starting agent_master_id=%s", idx, len(pending), agent_id)

            try:
                rows, attempts = fetch_listings(
                    conn,
                    agent_id,
                    query_timeout_sec=args.query_timeout,
                    max_retries=args.max_retries,
                )
                elapsed = time.perf_counter() - started
                if rows:
                    append_excel_rows(output, rows, RESULT_COLUMNS)
                append_progress(output, agent_id, "ok", len(rows))
                total_rows += len(rows)
                logger.info(
                    "[%s/%s] id=%s rows=%s (%.1fs)",
                    idx,
                    len(pending),
                    agent_id,
                    len(rows),
                    elapsed,
                )
            except RuntimeError as exc:
                errors += 1
                exit_code = 1
                error_type = getattr(exc, "error_type", "error")
                log_failure(
                    output,
                    error_logger,
                    agent_id,
                    error_type=error_type,
                    message=str(exc),
                    duration_sec=time.perf_counter() - started,
                    mysql_errno=getattr(exc, "mysql_errno", None),
                    attempts=getattr(exc, "attempts", 1),
                )
                if error_type == "connection_error":
                    conn = reconnect_db(cfg, conn, read_timeout=read_timeout)
            except (OSError, PermissionError) as exc:
                errors += 1
                exit_code = 1
                log_failure(
                    output,
                    error_logger,
                    agent_id,
                    error_type="excel_error",
                    message=str(exc),
                    duration_sec=time.perf_counter() - started,
                )
            except Exception as exc:
                errors += 1
                exit_code = 1
                log_failure(
                    output,
                    error_logger,
                    agent_id,
                    error_type="unknown_error",
                    message=f"{type(exc).__name__}: {exc}",
                    duration_sec=time.perf_counter() - started,
                )
                logger.exception("Unexpected error for id=%s", agent_id)

            if args.sleep > 0 and idx < len(pending) and not _shutdown_requested:
                time.sleep(args.sleep)

        if not args.skip_birds_eye and output.exists() and not interrupted:
            try:
                summary = add_birds_eye_sheet(output)
                totals = summary.iloc[-1]
                logger.info(
                    "Birds Eye View added: branches=%s total=%s active=%s sold/uo=%s withdrawn=%s",
                    len(summary) - 1,
                    int(totals["Total Listings"]),
                    int(totals["Active"]),
                    int(totals["Sold / Under Offer"]),
                    int(totals["Withdrawn"]),
                )
            except Exception as exc:
                exit_code = 1
                logger.error("Birds Eye View generation failed: %s", exc)

        logger.info(
            "Done. listing_rows=%s errors=%s interrupted=%s output=%s",
            total_rows,
            errors,
            interrupted,
            output,
        )
        return 130 if interrupted else exit_code

    except KeyboardInterrupt:
        logger.warning("KeyboardInterrupt — progress saved. Re-run with --resume.")
        return 130
    finally:
        try:
            conn.close()
        except Exception:
            pass


if __name__ == "__main__":
    raise SystemExit(main())
